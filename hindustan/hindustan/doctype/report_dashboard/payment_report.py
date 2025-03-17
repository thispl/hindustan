import frappe
from frappe.model.document import Document

import frappe
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Side,Font,PatternFill

@frappe.whitelist()
def payment_report_download():
    posting_date = datetime.now().strftime("%d-%m-%Y")
    filename_todo = "Payment Report"+ posting_date
    build_xlsx_response_todo(filename_todo)
    
def build_xlsx_response_todo(filename_todo):
    xlsx_file = make_xlsx_todo(filename_todo)
    frappe.response['filename'] = filename_todo + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
def make_xlsx_todo(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)
    fill_color = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)  # White color font
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 25 
    ws.column_dimensions['H'].width = 25 
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    black_border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )
    title_font = Font(bold=True, size=14)
    alignment = Alignment(horizontal="center")
    text_wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")  # Centered with wrap

    posting_date = datetime.now().strftime("%d-%m-%Y")
    # title = "TODO Report (" + posting_date + ")"
    # # ws.merge_cells("A1:F1")
    # ws["A1"].value = title
    # ws["A1"].font = title_font
    # ws["A1"].alignment = alignment
    bold_font = Font(bold=True)
    header=['Payment Report'+' '+str(posting_date)]
    ws.append(header)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    header_cell = ws.cell(row=1, column=1)
    header_cell.alignment = alignment
    header_cell.font = bold_font
    header = ["S NO", "Employee", "Employee Name", "Institution", "Department", "Bank",'Account Number','IFSC Code','Net Pay']
    ws.append(header) 
    for cell in ws[2]: 
        cell.font = header_font  # Apply white font to each header cell 
        cell.border = black_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = fill_color
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):  # Loop through the first column
        for cell in row:
            if cell.value == "Total":
                ws.merge_cells(start_row=cell.row, start_column=1, end_row=cell.row, end_column=8)
                cell.font = bold_font
    
    data1= get_data_of_todo(args)
    for row in data1:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = text_wrap
            cell.border = thin_border
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):  # Loop through the first column
        for cell in row:
            if cell.value == "Total":
                ws.merge_cells(start_row=cell.row, start_column=1, end_row=cell.row, end_column=8)
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def get_data_of_todo(args):
    data = []
    sdate=args.from_date
    tdate=args.to_date
    ins=args.institute_name
    dept=args.dept
    s_no=1
    net=0
    if ins and dept:
        slip = frappe.db.get_all("Salary Slip",{"start_date":sdate,'end_date':tdate,"docstatus":['!=',2],'custom_institute_name':ins,'department':dept}, ["*"])
    elif not ins and dept:
        slip = frappe.db.get_all("Salary Slip",{"start_date":sdate,'end_date':tdate,"docstatus":['!=',2],'department':dept}, ["*"])
    elif ins and not dept:
        slip = frappe.db.get_all("Salary Slip",{"start_date":sdate,'end_date':tdate,"docstatus":['!=',2],'custom_institute_name':ins}, ["*"])
    else:
        slip = frappe.db.get_all("Salary Slip",{"start_date":sdate,'end_date':tdate,"docstatus":['!=',2]}, ["*"])
    for i in slip:
        bank=frappe.db.get_value("Employee",{'name':i.employee},['bank_name'])
        acc=frappe.db.get_value("Employee",{'name':i.employee},['bank_ac_no'])
        ifsc=frappe.db.get_value("Employee",{'name':i.employee},['ifsc_code'])
        data.append([s_no, i.employee, i.employee_name, i.department,i.custom_institute_name, bank or '-',acc or '-',ifsc or '-',i.net_pay])
        net+=i.net_pay
        s_no+=1
    data.append(['Total','','','','','','','',net])
    return data
