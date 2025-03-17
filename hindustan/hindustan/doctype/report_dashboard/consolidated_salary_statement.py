import frappe
from frappe.model.document import Document


import frappe
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Side,Font

@frappe.whitelist()
def consolidated_salary_statement_download():
    filename = "CONSOLIDATED SALARY STATEMENT"
    build_xlsx_response_css(filename)

def build_xlsx_response_css(filename):
    xlsx_file = make_xlsx_css(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def make_xlsx_css(sheet_name="CONSOLIDATED SALARY STATEMENT", wb=None):
    if wb is None:
        wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = sheet_name
    args = frappe.local.form_dict

    from_date_str = args.get('from_date')
    to_date_str = args.get('to_date')
    institute=args.get('institute_name')
    if not from_date_str or not to_date_str:
        frappe.throw("Both 'From date' and 'To date' must be provided.")

    from_date = datetime.strptime(from_date_str, '%Y-%m-%d')
    to_date = datetime.strptime(to_date_str, '%Y-%m-%d')

    # Define column widths
    column_widths = {
        'A': 5, 'B': 20, 'C': 20, 'D': 20, 'E': 20, 'F': 20
    
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Adding header rows
    if institute:
        ws.append([f"{institute}"])
    else:
        ws.append(["HINDUSTAN ACADEMY"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')

    ws.append(["BANGALORE"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal='center', vertical='center')

    month_year = from_date.strftime('%B %Y')
    month_year=str(month_year).upper()
    ws.append([f"CONSOLIDATED SALARY STATEMENT FOR THE MONTH - {month_year}"])
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    ws.cell(row=3, column=1).alignment = Alignment(horizontal='center', vertical='center')

    # Adding sub-header
    sub_header = [
        "SI NO", "Section", "Directly to Bank", "By Cheque", "By Cash", "Total"
    ]
    ws.append(sub_header)
    bold_font = Font(bold=True)
    for cell in ws[4]:
        cell.font = bold_font
    # Define border style
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Fetch data
    data = get_data_for_css(args)

    # Populate data rows
    s_no = 1
    totals = [0] *4   # Initialize totals for numeric columns
    for row in data:
        row_data = [
            s_no,
            row.get("department", "-"),
            row.get("bank_amount", 0),
            row.get("cheque_amount", 0),
            row.get("cash_amount", 0),
            row.get("total_amount", 0),
        ]
        ws.append(row_data)

        # Accumulate totals for numeric columns
        totals[0] += row.get("bank_amount", 0) or 0
        totals[1] += row.get("cheque_amount", 0) or 0
        totals[2] += row.get("cash_amount", 0) or 0
        totals[3] += row.get("total_amount", 0) or 0
        s_no += 1

    # Add totals row
    totals_row = [
        "TOTAL", "-",totals[0],totals[1],totals[2],totals[3]
    ]
    ws.append(totals_row)
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
    total_cell = ws.cell(row=ws.max_row, column=1)  # This is the merged cell (SI NO)
    total_cell.alignment = Alignment(horizontal='center', vertical='center')
    for cell in ws[ws.max_row]:  # Total row
        cell.font = bold_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    # Apply borders to all rows and columns
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
    font_size=14
    for row in ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=19):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = Font(size=font_size)
    # Save the workbook to a BytesIO object
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

    return xlsx_file

def get_data_for_css(args):
    data={}
    s_no=1
    from_date = args.get("from_date")
    to_date = args.get("to_date")
    institute = args.get("institute_name") 
    dept = args.get("dept") 
    if not from_date or not to_date:
        frappe.throw("Both 'From date' and 'To date' are required.")

    institute_condition = ""
    if institute and not dept:
        institute_condition = "AND e.custom_institute_name = %s"
        params = (from_date, to_date, institute)
    elif not institute and dept:
        institute_condition = "AND e.department = %s"
        params = (from_date, to_date, dept)
    elif institute and dept:
        institute_condition = "AND e.custom_institute_name = %s AND e.department = %s"
        params = (from_date, to_date, institute, dept)
    else:
        institute_condition = ""
        params = (from_date, to_date)
    bank = frappe.db.sql("""
        SELECT 
            e.department AS department, 
            SUM(s.net_pay) AS bank_amount
        FROM 
            `tabSalary Slip` s
        LEFT JOIN 
            `tabEmployee` e ON s.employee = e.name
        WHERE 
            s.start_date BETWEEN %s AND %s
            AND e.salary_mode='Bank'
            {institute_condition}
        GROUP BY 
            e.department
    """.format(institute_condition=institute_condition), 
    params, as_dict=True)
    cheque = frappe.db.sql("""
        SELECT 
            e.department AS department, 
            SUM(s.net_pay) AS cheque_amount
        FROM 
            `tabSalary Slip` s
        LEFT JOIN 
            `tabEmployee` e ON s.employee = e.name
        WHERE 
            s.start_date BETWEEN %s AND %s
            AND e.salary_mode='Cheque'
            {institute_condition}
        GROUP BY 
            e.department
    """.format(institute_condition=institute_condition), 
    params, as_dict=True)
    cash = frappe.db.sql("""
        SELECT 
            e.department AS department, 
            SUM(s.net_pay) AS cash_amount
        FROM 
            `tabSalary Slip` s
        LEFT JOIN 
            `tabEmployee` e ON s.employee = e.name
        WHERE 
            s.start_date BETWEEN %s AND %s
            AND e.salary_mode='Cash'
            {institute_condition}
        GROUP BY 
            e.department
    """.format(institute_condition=institute_condition), 
    params, as_dict=True)
    total = frappe.db.sql("""
        SELECT 
            e.department AS department, 
            SUM(s.net_pay) AS total_amount
        FROM 
            `tabSalary Slip` s
        LEFT JOIN 
            `tabEmployee` e ON s.employee = e.name
        WHERE 
            s.start_date BETWEEN %s AND %s
            AND e.salary_mode IN ('Cash','Cheque','Bank')
            {institute_condition}
        GROUP BY 
            e.department
    """.format(institute_condition=institute_condition), 
    params, as_dict=True)
    for item in bank:
        department = item['department']
        data[department] = {
            "department": department,
            "bank_amount": item['bank_amount'],
            "cheque_amount": 0,
            "cash_amount": 0,
            "total_amount": 0,
        }

    for item in cheque:
        department = item['department']
        if department not in data:
            data[department] = {
                "department": department,
                "bank_amount": 0,
                "cheque_amount": item['cheque_amount'],
                "cash_amount": 0,
                "total_amount": 0,
            }
        else:
            data[department]['cheque_amount'] = item['cheque_amount']

    for item in cash:
        department = item['department']
        if department not in data:
            data[department] = {
                "department": department,
                "bank_amount": 0,
                "cheque_amount": 0,
                "cash_amount": item['cash_amount'],
                "total_amount": 0,
            }
        else:
            data[department]['cash_amount'] = item['cash_amount']

    for item in total:
        department = item['department']
        if department not in data:
            data[department] = {
                "department": department,
                "bank_amount": 0,
                "cheque_amount": 0,
                "cash_amount": 0,
                "total_amount": item['total_amount'],
            }
        else:
            data[department]['total_amount'] = item['total_amount']
    return list(data.values())


@frappe.whitelist()
def print_consolidated_salary(doc):
    from datetime import datetime

    from_date = doc.from_date
    to_date = doc.to_date
    institute = doc.institute_name
    dept = doc.department

    if not from_date or not to_date:
        frappe.throw("Both 'From date' and 'To date' are required.")

    formatted_date = datetime.strptime(from_date, "%Y-%m-%d").strftime("%Y")
    formatted_month = (datetime.strptime(from_date, "%Y-%m-%d").strftime("%B")).upper()
    formated=str(formatted_month)+' '+str(formatted_date)
    args = {"from_date": from_date, "to_date": to_date, "institute_name": institute, 'dept':dept}
    salary_data = get_data_for_css(args)

    total_bank_amount = 0
    total_cheque_amount = 0
    total_cash_amount = 0
    grand_total = 0

    data = '<table style="border-collapse: collapse; width: 100%; border: 1px solid black;">'
    data += f'<tr><td colspan="12" style="text-align:center; font-weight:bold; border: 1px solid black;font-size:14px;">CONSOLIDATED SALARY STATEMENT FOR THE MONTH - {formated}</td></tr>'
    data += f'<tr><td colspan="12" style="text-align:center; font-weight:bold; border: 1px solid black;font-size:14px;">BANGALORE</td></tr>'
    data += '<tr>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">S.NO</td>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">Section</td>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">Directly to Bank</td>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">By Cheque</td>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">By Cash</td>' \
            '<td style="text-align:center; font-weight:bold; border: 1px solid black;">Total</td>' \
            '</tr>'

    s_no = 1
    for record in salary_data:
        bank_amount = record["bank_amount"] or 0
        cheque_amount = record["cheque_amount"] or 0
        cash_amount = record["cash_amount"] or 0
        total_amount = record["total_amount"] or 0

        total_bank_amount += bank_amount
        total_cheque_amount += cheque_amount
        total_cash_amount += cash_amount
        grand_total += total_amount

        data += f'<tr>' \
                f'<td style="text-align:center; border: 1px solid black;">{s_no}</td>' \
                f'<td style="text-align:center; border: 1px solid black;">{record["department"]}</td>' \
                f'<td style="text-align:center; border: 1px solid black;">{bank_amount}</td>' \
                f'<td style="text-align:center; border: 1px solid black;">{cheque_amount}</td>' \
                f'<td style="text-align:center; border: 1px solid black;">{cash_amount}</td>' \
                f'<td style="text-align:center; border: 1px solid black;">{total_amount}</td>' \
                '</tr>'
        s_no += 1

    data += f'<tr style="font-weight:bold;">' \
            f'<td colspan="2" style="text-align:center; border: 1px solid black;">Total</td>' \
            f'<td style="text-align:center; border: 1px solid black;">{total_bank_amount}</td>' \
            f'<td style="text-align:center; border: 1px solid black;">{total_cheque_amount}</td>' \
            f'<td style="text-align:center; border: 1px solid black;">{total_cash_amount}</td>' \
            f'<td style="text-align:center; border: 1px solid black;">{grand_total}</td>' \
            '</tr>'

    data += '</table>'

    return data