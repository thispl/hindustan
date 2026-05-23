# # Copyright (c) 2024, Abdulla PI and contributors
# # For license information, please see license.txt

import frappe
from frappe.model.document import Document
import math


class ReportDashboard(Document):
    pass

import frappe
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Side,Font
from frappe.utils.pdf import get_pdf
from frappe.utils.response import build_response

@frappe.whitelist()
def transfer_statement_download():
    filename = "TRANSFER STATEMENT"
    build_xlsx_response_ts(filename)

def build_xlsx_response_ts(filename):
    xlsx_file = make_xlsx_css(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def make_xlsx_css(sheet_name="TRANSFER STATEMENT", wb=None):
    if wb is None:
        wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = sheet_name
    args = frappe.local.form_dict
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    from_date_str = args.get('from_date')
    to_date_str = args.get('to_date')
    institute=args.get('institute_name')
    if not from_date_str or not to_date_str:
        frappe.throw("Both 'From date' and 'To date' must be provided.")

    from_date = datetime.strptime(from_date_str, '%Y-%m-%d')
    to_date = datetime.strptime(to_date_str, '%Y-%m-%d')

    # Define column widths
    column_widths = {
        'A': 7, 'B': 18, 'C': 25, 'D': 25,'E':25,'F':25,'G':10
    
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Adding header rows
    if institute:
        ws.append([f"{institute}"])
    else:
        ws.append(["HINDUSTAN ACADEMY"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')

    ws.append(["BANGALORE"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal='center', vertical='center')

    section_name = args.get("dept")
    if section_name:
        section_display = f"SECTION : {section_name.upper()}"
    else:
        section_display = "SECTION : ALL SECTIONS"
    ws.append([section_display])
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
    ws.cell(row=3, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=3, column=1).font = Font(bold=False)
    
    month_year = from_date.strftime('%B %Y')
    month_year=str(month_year).upper()
    ws.append([f"TRANSFER STATEMENT FOR THE MONTH - {month_year}"])
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=7)
    ws.cell(row=3, column=1).alignment = Alignment(horizontal='center', vertical='center')
    
    sub_header = [
        "SI NO", "Employee","Employee Name" ,"Designation", "Bank Name", "Account Number", "Net Pay"
    ]
    ws.append(sub_header)
    bold_font = Font(bold=True)
    for cell in ws[5]:
        cell.font = bold_font
    data = get_data_for_ts(args)
    total_net_pay = 0
    for row in data:
        ws.append(row)
        total_net_pay += row[6]
    total_row = ["TOTAL", "", "", "", "", "", total_net_pay]
    ws.append(total_row)

    last_row = ws.max_row

    ws.cell(row=last_row, column=7).font = Font(bold=True)
    ws.cell(row=last_row, column=7).font = Font(bold=True)

    ws.cell(row=last_row, column=7).alignment = Alignment(horizontal='right')
    ws.cell(row=last_row, column=7).alignment = Alignment(horizontal='center')
    ws.merge_cells(start_row=last_row, start_column=1, end_row=last_row, end_column=6)

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # Save the workbook to a BytesIO object
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

    return xlsx_file

# def get_data_for_ts(args):
#     s_no = 0
#     data = []
#     filters = {"start_date": args.from_date, "docstatus": ['!=', 2]}
#     if args.dept and args.institute:   
#         salary_slips = frappe.db.get_all("Salary Slip", {"start_date": args.from_date,'custom_institute_name':args.institute,'department':args.dept,'docstatus':['!=',2]}, ['employee', 'employee_name', 'net_pay'])
#     elif args.dept and not args.institute:   
#         salary_slips = frappe.db.get_all("Salary Slip", {"start_date": args.from_date,'department':args.dept,'docstatus':['!=',2]}, ['employee', 'employee_name', 'net_pay'])
#     elif not args.dept and args.institute:   
#         salary_slips = frappe.db.get_all("Salary Slip", {"start_date": args.from_date,'custom_institute_name':args.institute,'docstatus':['!=',2]}, ['employee', 'employee_name', 'net_pay'])
#     else:   
#         salary_slips = frappe.db.get_all("Salary Slip", {"start_date": args.from_date,'docstatus':['!=',2]}, ['employee', 'employee_name', 'net_pay'])
#     employees = frappe.db.get_all("Employee", 
#                                   {"name": ["in", [slip['employee'] for slip in salary_slips]]},
#                                   ["name", "bank_ac_no", "ifsc_code", "bank_name"])
    
#     employee_data = {emp['name']: emp for emp in employees}
#     for slip in salary_slips:
#         s_no += 1
#         emp_details = employee_data.get(slip['employee'], {})
#         bank_name = emp_details.get("bank_name", '')
#         bank_account = emp_details.get("bank_ac_no", '')
#         ifsc_code = emp_details.get("ifsc_code", '')
#         row = [s_no, slip['employee'], slip['employee_name'], bank_name or '-', bank_account or '-', ifsc_code or '-', slip['net_pay']]
#         data.append(row)
    
#     return data
 

    
def get_data_for_ts(args):
    s_no = 0
    data = []
    from_date = args.get("from_date")
    institute = args.get("institute_name")
    dept = args.get("dept")
    employee_filters = {}
    if institute:
        employee_filters["custom_institute_name"] = institute

    if dept:
        employee_filters["custom_section"] = dept

    employees = frappe.db.get_all("Employee",employee_filters, ["name", "employee_name", "bank_ac_no", "bank_name"])
    if not employees:
        return []
    employee_map = {e.name: e for e in employees}

    salary_slips = frappe.db.get_all( "Salary Slip", {"start_date": from_date,"employee": ["in", list(employee_map.keys())],"docstatus": ["!=", 2] },["employee", "employee_name", "designation","net_pay"])

    designation_list = list(set([s.designation for s in salary_slips if s.designation]))

    designations = frappe.db.get_all(
        "Designation",
        {"name": ["in", designation_list]},
        [
            "name",
            "custom_office_order_no",
            "custom_haa_order_no",
            "custom_haeas_order_no",
            "custom_hea_order_no"
        ]
    )

    designation_map = {d.name: d for d in designations}

    temp_data = []

    for slip in salary_slips:
        emp = employee_map.get(slip.employee)
        desig = designation_map.get(slip.designation)

        if not emp or not desig:
            continue

        section = emp.custom_section

        # Decide order field
        if dept:
            if dept == "Office":
                order_no = desig.custom_office_order_no or 0
            elif dept == "HAA":
                order_no = desig.custom_haa_order_no or 0
            elif dept == "HBS":
                order_no = desig.custom_haeas_order_no or 0
            elif dept == "HEA":
                order_no = desig.custom_hea_order_no or 0
            else:
                order_no = 0
        else:
            # No filter → common priority
            if section == "Office":
                order_no = desig.custom_office_order_no or 0
            elif section == "HAA":
                order_no = desig.custom_haa_order_no or 0
            elif section == "HBS":
                order_no = desig.custom_haeas_order_no or 0
            elif section == "HEA":
                order_no = desig.custom_hea_order_no or 0
            else:
                order_no = 0

        net_pay_rounded = math.floor((slip.net_pay or 0) + 0.5)
        order_no_for_sort = order_no if order_no != 0 else 9999
        temp_data.append({
            "order_no": order_no_for_sort,  
            "section": section,
            "row": [
                slip.employee,
                slip.employee_name,
                slip.designation,
                emp.bank_name or '-',
                emp.bank_ac_no or '-',
                net_pay_rounded
            ]
        })

    temp_data.sort(key=lambda x: x["order_no"])

    for item in temp_data:
        s_no += 1
        data.append([s_no] + item["row"])

    return data


@frappe.whitelist()
def transfer_statement_pdf_download():
    args = frappe.local.form_dict
    data = get_data_for_ts(args)

    institute = args.get("institute_name") or "HINDUSTAN ACADEMY"
    dept = args.get("dept") or "ALL SECTIONS"

    total_net_pay = sum([row[6] for row in data])

    html = f"""
        <style>

            body {{
                font-family: Arial, sans-serif;
            }}

            .title {{
                text-align:center;
                font-weight:bold;
            }}

            .title.main {{
                font-size:20px;
            }}

            .title.sub {{
                font-size:14px;
            }}

            table {{
                width:100%;
                border-collapse:collapse;
                font-size:12px;  
            }}

            th {{
                font-size:12px;   
                padding:6px;
            }}

            td {{
                font-size:12px;
                padding:5px;
            }}

            th, td {{
                border:1px solid black;
                text-align:center;
            }}

        </style>

        <div class="title main">Evehans Academy</div>
        <div class="title sub">BANGALORE</div>
        <div class="title sub">SECTION : {dept}</div>
        <div class="title sub">TRANSFER STATEMENT</div>
        <br>

    <table>
        <thead>
            <tr>
                <th>SI NO</th>
                <th>Employee</th>
                <th>Employee Name</th>
                <th>Designation</th>
                <th>Bank Name</th>
                <th>Account Number</th>
                <th>Net Pay</th>
            </tr>
        </thead>
        <tbody>
    """

    for row in data:
        html += f"""
        <tr>
            <td>{row[0]}</td>
            <td>{row[1]}</td>
            <td>{row[2]}</td>
            <td>{row[3]}</td>
            <td>{row[4]}</td>
            <td>{row[5]}</td>
            <td>{row[6]}</td>
        </tr>
        """

    html += f"""
        <tr>
            <td colspan="6"><b>TOTAL</b></td>
            <td><b>{total_net_pay}</b></td>
        </tr>
        </tbody>
    </table>
    """

    options = {
        "orientation": "Landscape",
        "page-size": "A4",
        "margin-top": "10mm",
        "margin-bottom": "10mm",
        "margin-left": "5mm",
        "margin-right": "5mm"
    }

    pdf = get_pdf(html, options)

    frappe.response.filename = "Transfer Statement.pdf"
    frappe.response.filecontent = pdf
    frappe.response.type = "binary"